#!/usr/bin/env python3
"""Rebuild every derived inventory file from the source workbooks.

Current list    : Data/source/<SharePoint export>.xlsx  (one flat sheet, the
                  whole set as exported from the SharePoint kit catalogue)
Previous list   : Data/legacy/*.xlsx  (one workbook per kit, the list the app
                  was first built from)
Generated       : Data/inventory.json
                  Data/kit_items.csv
                  Data/master_inventory.csv
                  Data/kits/<slug>.csv

The SharePoint export is authoritative for every item it lists: NSN, quantity,
unit of measure, description and price. Items that appear only in the previous
list are kept, marked inactive, so the history is not lost.

Run after replacing the export:

    python3 tools/build_inventory.py

Requires openpyxl (pip install openpyxl).
"""

import csv
import json
import os
import re
import sys
import unicodedata
from collections import OrderedDict

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOURCE_DIR = os.path.join(ROOT, "Data", "source")
LEGACY_DIR = os.path.join(ROOT, "Data", "legacy")
DATA_DIR = os.path.join(ROOT, "Data")
KIT_DIR = os.path.join(DATA_DIR, "kits")

CURRENT = "current"
INACTIVE = "inactive"

# Column order in every legacy workbook.
LEGACY_COLUMNS = ["kit_nsn", "item_nsn", "qty", "uom", "description",
                  "source", "price", "accountability_code"]

LEGACY_HEADER = ["kit nsn", "kit item nsn", "qty", "uom", "description",
                 "source", "price", "d kit accountability code"]

# SharePoint export columns this build reads, by the English half of their
# bilingual header ("Item NSN / NNO de l'article" -> "item nsn"). Every other
# column in the export is ignored.
SHAREPOINT_COLUMNS = {
    "kit nsn": "set_nsn",
    "kit name": "set_name",
    "item nsn": "item_nsn",
    "item name": "description",
    "total qty": "total_qty",
    "uom": "uom",
    "total value": "total_value",
    "qty per sub-kit": "qty_per_sub_kit",
    "# sub-kits req": "sub_kits_req",
    "sub-kit nsn": "sub_kit_nsn",
    "sub-kit name": "sub_kit_name",
    "full sub-kit path": "path",
}

NSN_RE = re.compile(r"^\d{4}-[0-9A-Z]{2}-[0-9A-Z]{3}-[0-9A-Z]{4}$")
LEADING_NUMBER_RE = re.compile(r"^\s*([0-9]*\.?[0-9]+)")


def slugify(value):
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    value = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-").lower()
    return re.sub(r"-{2,}", "-", value)


def short_slug(value, limit=50):
    """A slug cut at a word boundary, for kits named only by a long description."""
    slug = slugify(value)
    if len(slug) <= limit:
        return slug
    return slug[:limit + 1].rsplit("-", 1)[0]


def clean(value):
    """Normalise a cell to a trimmed string ('' for blanks)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def parse_number(value):
    """Parse a quantity/price cell. Returns None when not numeric."""
    text = clean(value).replace(",", "").replace("$", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def xlsx_files(directory):
    files = sorted(f for f in os.listdir(directory) if f.endswith(".xlsx")
                   and not f.startswith("~$"))
    if not files:
        sys.exit("no workbooks found in %s" % directory)
    return files


def single_sheet(path, filename):
    book = openpyxl.load_workbook(path, data_only=True)
    if len(book.worksheets) > 1:
        sys.exit("%s: expected one sheet, found %s" % (filename, book.sheetnames))
    return book.worksheets[0]


# ── Previous list ──────────────────────────────────────────────────────────

def read_legacy_rows():
    """Read every legacy per-kit workbook into a flat list of line dicts."""
    rows = []
    for filename in xlsx_files(LEGACY_DIR):
        sheet = single_sheet(os.path.join(LEGACY_DIR, filename), filename)
        raw = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        got = [clean(cell).lower() for cell in raw][:len(LEGACY_HEADER)]
        if got != LEGACY_HEADER:
            # Cells map to columns by position, so a reordered workbook would be
            # parsed wrongly and silently. Refuse to guess.
            sys.exit("%s: unexpected column layout.\n  expected: %s\n  found:    %s"
                     % (filename, LEGACY_HEADER, got))
        for excel_row, raw in enumerate(sheet.iter_rows(min_row=2, values_only=True),
                                        start=2):
            values = [clean(cell) for cell in raw][:len(LEGACY_COLUMNS)]
            values += [""] * (len(LEGACY_COLUMNS) - len(values))
            row = dict(zip(LEGACY_COLUMNS, values))
            if not row["kit_nsn"] and not row["item_nsn"]:
                continue
            # NSNs are case-insensitive identifiers; a few cells arrive lowercase.
            row["kit_nsn"] = row["kit_nsn"].upper()
            row["item_nsn"] = row["item_nsn"].upper()
            row["uom"] = row["uom"].upper()
            row["accountability_code"] = row["accountability_code"].upper()
            row["qty"] = parse_number(row["qty"])
            row["price"] = parse_number(row["price"])
            row["source_file"] = "legacy/" + filename
            row["source_row"] = excel_row
            rows.append(row)
    return rows


# ── Current list (SharePoint export) ───────────────────────────────────────

def header_key(cell):
    """'Item NSN / NNO de l'article' -> 'item nsn'."""
    return re.split(r"\s+[/|]\s+", clean(cell), maxsplit=1)[0].strip().lower()


def read_sharepoint_rows():
    files = xlsx_files(SOURCE_DIR)
    if len(files) != 1:
        sys.exit("expected exactly one SharePoint export in %s, found %s"
                 % (SOURCE_DIR, files))
    filename = files[0]
    sheet = single_sheet(os.path.join(SOURCE_DIR, filename), filename)

    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    index = {}
    for position, cell in enumerate(header):
        key = header_key(cell)
        if key in SHAREPOINT_COLUMNS and key not in index:
            index[key] = position
    missing = sorted(set(SHAREPOINT_COLUMNS) - set(index))
    if missing:
        sys.exit("%s: missing expected columns: %s" % (filename, missing))

    rows = []
    for excel_row, raw in enumerate(sheet.iter_rows(min_row=2, values_only=True),
                                    start=2):
        row = {field: raw[index[key]] if index[key] < len(raw) else None
               for key, field in SHAREPOINT_COLUMNS.items()}
        # Blank spacer rows and the "Applied filters" footer have no item.
        if not clean(row["item_nsn"]):
            continue
        for field in ("set_nsn", "set_name", "item_nsn", "description", "uom",
                      "sub_kit_nsn", "sub_kit_name", "path"):
            row[field] = clean(row[field])
        row["set_nsn"] = row["set_nsn"].upper()
        row["item_nsn"] = row["item_nsn"].upper()
        row["sub_kit_nsn"] = row["sub_kit_nsn"].upper()
        row["uom"] = row["uom"].upper()
        row["total_qty"] = parse_number(row["total_qty"])
        row["total_value"] = parse_number(row["total_value"])
        row["sub_kits_req"] = parse_number(row["sub_kits_req"])
        if row["total_qty"] is None:
            sys.exit("%s row %d: no Total Qty for %s" % (filename, excel_row,
                                                         row["item_nsn"]))
        # "Qty per Sub-Kit" arrives as text such as "70 EA" and is filled in only
        # when the sub-kit is held more than once. Otherwise Total Qty is already
        # the quantity in the one sub-kit.
        per = LEADING_NUMBER_RE.match(clean(row["qty_per_sub_kit"]))
        if per:
            row["qty"] = float(per.group(1))
        else:
            row["qty"] = row["total_qty"] / (row["sub_kits_req"] or 1)
        row["source_file"] = filename
        row["source_row"] = excel_row
        rows.append(row)
    if not rows:
        sys.exit("%s: no item rows found" % filename)
    return filename, rows


def path_ancestors(path):
    """'A ➡️ B | NAME ➡️ C' -> ['A', 'B']: the kits above the row's sub-kit."""
    head = path.split("|", 1)[0]
    return [part.strip().upper() for part in head.split("➡️")
            if NSN_RE.match(part.strip().upper())]


def path_last_name(path):
    """The (possibly truncated) name of the sub-kit's parent, from the path."""
    if "|" not in path:
        return ""
    return path.split("|", 1)[1].split("➡️")[0].strip().rstrip(".").strip()


# ── Merge ──────────────────────────────────────────────────────────────────

def build(legacy_rows, sp_file, sp_rows):
    sets = {r["set_nsn"] for r in sp_rows if r["set_nsn"]}
    if len(sets) != 1:
        sys.exit("expected one top-level set in the export, found %s" % sorted(sets))
    root = sets.pop()
    root_name = next(r["set_name"] for r in sp_rows if r["set_name"])

    # Legacy lookups: a kit's issued name is its description where it appears
    # as a component; source / accountability code carry across by NSN.
    legacy_kits = OrderedDict()
    for r in legacy_rows:
        legacy_kits.setdefault(r["kit_nsn"], r["source_file"])
    legacy_name, legacy_parent = {}, {}
    legacy_line, legacy_item = {}, {}
    for r in legacy_rows:
        if r["item_nsn"] in legacy_kits:
            legacy_name.setdefault(r["item_nsn"], r["description"])
            legacy_parent.setdefault(r["item_nsn"], r["kit_nsn"])
        legacy_line[(r["kit_nsn"], r["item_nsn"])] = r
        legacy_item.setdefault(r["item_nsn"], r)

    # 1. Current kits and their hierarchy, from each row's sub-kit path.
    kit_name = {root: root_name}
    kit_parent = {root: None}
    kit_mult = {root: 1.0}
    kit_order = [root]
    path_names = {}
    for r in sp_rows:
        nsn = r["sub_kit_nsn"]
        if not nsn:
            continue
        chain = path_ancestors(r["path"]) + [nsn]
        if chain[0] != root:
            sys.exit("%s row %d: path does not start at %s: %s"
                     % (sp_file, r["source_row"], root, r["path"]))
        for parent, child in zip(chain, chain[1:]):
            if kit_parent.setdefault(child, parent) != parent:
                sys.exit("%s: kit %s has two parents (%s, %s)"
                         % (sp_file, child, kit_parent[child], parent))
            if child not in kit_order:
                kit_order.append(child)
        if len(chain) > 2:
            path_names.setdefault(chain[-2], path_last_name(r["path"]))
        kit_name.setdefault(nsn, r["sub_kit_name"])
        if r["sub_kits_req"] is not None:
            if kit_mult.setdefault(nsn, r["sub_kits_req"]) != r["sub_kits_req"]:
                sys.exit("%s: kit %s has inconsistent '# Sub-Kits Req'" % (sp_file, nsn))
    for nsn in kit_order:
        # A kit that only ever appears as an intermediate step in a path (the
        # core treatment set) has no name column of its own; the legacy list's
        # name is complete, the path's is truncated.
        if not kit_name.get(nsn):
            kit_name[nsn] = legacy_name.get(nsn) or path_names.get(nsn) or nsn
        kit_mult.setdefault(nsn, 1.0)

    # 2. Current lines: every export row, plus one container line per sub-kit
    #    inside its parent (the export omits kit definitions as items).
    lines = OrderedDict()
    for r in sp_rows:
        kit = r["sub_kit_nsn"] or root
        key = (kit, r["item_nsn"])
        if key in lines:
            sys.exit("%s row %d: %s listed twice in %s"
                     % (sp_file, r["source_row"], r["item_nsn"], kit))
        lines[key] = {
            "kit_nsn": kit, "item_nsn": r["item_nsn"], "qty": r["qty"],
            "uom": r["uom"], "description": r["description"],
            "status": CURRENT, "source_file": r["source_file"],
            "source_row": r["source_row"],
        }
    for nsn in kit_order[1:]:
        parent = kit_parent[nsn]
        key = (parent, nsn)
        if key in lines:
            continue
        lines[key] = {
            "kit_nsn": parent, "item_nsn": nsn,
            "qty": kit_mult[nsn] / kit_mult[parent],
            "uom": "KT", "description": kit_name[nsn],
            "status": CURRENT, "source_file": sp_file, "source_row": None,
        }
    current_items = {k[1] for k in lines}

    # 3. Prices: SharePoint gives a total value per row, so the unit price is
    #    value / quantity, pooled across rows so rounding cannot make one item
    #    look differently priced in two kits. Where SharePoint shows no value
    #    but the previous list had a price, the previous price is kept.
    value_sum, qty_sum = {}, {}
    for r in sp_rows:
        value_sum[r["item_nsn"]] = value_sum.get(r["item_nsn"], 0) + (r["total_value"] or 0)
        qty_sum[r["item_nsn"]] = qty_sum.get(r["item_nsn"], 0) + r["total_qty"]
    unit_price, price_fallbacks = {}, []
    for nsn, value in value_sum.items():
        price = round(value / qty_sum[nsn], 2) if qty_sum[nsn] else None
        old = legacy_item.get(nsn)
        if not price and old and old["price"]:
            price = old["price"]
            price_fallbacks.append(nsn)
        unit_price[nsn] = price

    for line in lines.values():
        old = legacy_line.get((line["kit_nsn"], line["item_nsn"])) or \
            legacy_item.get(line["item_nsn"])
        line["source"] = old["source"] if old else ""
        line["accountability_code"] = old["accountability_code"] if old else ""
        line["price"] = unit_price.get(line["item_nsn"])

    # A container's price is the value of its contents, so price kits bottom-up.
    kit_value = {}

    def value_of(nsn, seen=()):
        if nsn in kit_value:
            return kit_value[nsn]
        if nsn in seen:
            sys.exit("kit hierarchy loops through %s" % nsn)
        total = 0.0
        for line in lines.values():
            if line["kit_nsn"] != nsn:
                continue
            if line["item_nsn"] in kit_name:
                line["price"] = round(value_of(line["item_nsn"], seen + (nsn,)), 2)
            if line["price"] is not None and line["qty"] is not None:
                total += line["qty"] * line["price"]
        kit_value[nsn] = round(total, 2)
        return kit_value[nsn]

    value_of(root)

    # 4. Inactive lines: legacy lines for items whose NSN is not on the current
    #    list at all. An item still current elsewhere is not carried into a kit
    #    it was dropped from.
    inactive_kits = OrderedDict()
    legacy_kit_lines = {}
    for r in legacy_rows:
        legacy_kit_lines[r["kit_nsn"]] = legacy_kit_lines.get(r["kit_nsn"], 0) + 1
    for r in legacy_rows:
        if r["item_nsn"] in current_items:
            continue
        key = (r["kit_nsn"], r["item_nsn"])
        lines[key] = {
            "kit_nsn": r["kit_nsn"], "item_nsn": r["item_nsn"], "qty": r["qty"],
            "uom": r["uom"], "description": r["description"],
            "source": r["source"], "price": r["price"],
            "accountability_code": r["accountability_code"],
            "status": INACTIVE, "source_file": r["source_file"],
            "source_row": r["source_row"],
        }
        if r["kit_nsn"] not in kit_name:
            inactive_kits.setdefault(r["kit_nsn"], legacy_kits[r["kit_nsn"]])

    for line in lines.values():
        line["is_nsn"] = bool(NSN_RE.match(line["item_nsn"]))
        line["extended_price"] = (round(line["qty"] * line["price"], 2)
                                  if line["qty"] is not None
                                  and line["price"] is not None else None)

    # 5. Kit records.
    legacy_slugs = {nsn: slugify(os.path.splitext(os.path.basename(f))[0])
                    for nsn, f in legacy_kits.items()}
    kits = OrderedDict()
    used_slugs = {}

    def add_kit(nsn, name, parent, status):
        # Keep the slug a kit already had, so existing CSV links keep working.
        slug = legacy_slugs.get(nsn) or short_slug(name) or nsn.lower()
        if used_slugs.get(slug, nsn) != nsn:
            slug = "%s-%s" % (slug, nsn.lower())
        used_slugs[slug] = nsn
        own = [l for l in lines.values() if l["kit_nsn"] == nsn]
        live = [l for l in own if l["status"] == CURRENT]
        kits[nsn] = {
            "nsn": nsn,
            "name": name,
            "slug": slug,
            "parent_nsn": parent,
            "status": status,
            "line_count": len(live),
            "total_value": round(sum(l["extended_price"] or 0 for l in live), 2),
            "inactive_line_count": len(own) - len(live),
            # Current sub-kits only; an inactive sub-kit is an inactive line.
            "sub_kits": [l["item_nsn"] for l in live if l["item_nsn"] in kit_name],
        }
        if status == INACTIVE:
            # How many lines the kit had in the previous list, so the app can say
            # how many are still current and appear in their current kits instead.
            kits[nsn]["former_line_count"] = legacy_kit_lines.get(nsn, 0)

    for nsn in kit_order:
        add_kit(nsn, kit_name[nsn], kit_parent[nsn], CURRENT)
    for nsn in inactive_kits:
        add_kit(nsn, legacy_name.get(nsn) or nsn, legacy_parent.get(nsn), INACTIVE)

    def depth_of(nsn, seen=()):
        parent = kits[nsn]["parent_nsn"]
        if not parent or parent not in kits or parent in seen:
            return 0
        return 1 + depth_of(parent, seen + (nsn,))

    for nsn in kits:
        kits[nsn]["depth"] = depth_of(nsn)

    stats = {"price_fallbacks": price_fallbacks}
    return kits, list(lines.values()), stats


def build_items(rows, kits):
    """Collapse the line rows into one record per distinct item."""
    items = OrderedDict()
    for row in rows:
        item = items.get(row["item_nsn"])
        if item is None:
            item = items[row["item_nsn"]] = {
                "nsn": row["item_nsn"],
                "description": row["description"],
                "uom": row["uom"],
                "source": row["source"],
                "accountability_code": row["accountability_code"],
                "status": row["status"],
                "is_nsn": row["is_nsn"],
                "is_kit": row["item_nsn"] in kits,
                "prices": [],
                "sources": [],
                "memberships": [],
            }
        elif item["status"] != row["status"]:
            sys.exit("%s is both current and inactive" % row["item_nsn"])
        if row["price"] is not None:
            item["prices"].append(row["price"])
        item["sources"].append(row["source"])
        item["memberships"].append({
            "kit_nsn": row["kit_nsn"],
            "qty": row["qty"],
            "price": row["price"],
            "extended_price": row["extended_price"],
        })

    for item in items.values():
        prices = item.pop("prices")
        # A few items differ between kits on price or on where they are drawn
        # from. The rollup has to pick one, so flag that it did.
        item["price"] = round(prices[0], 2) if prices else None
        item["price_varies"] = len(set(prices)) > 1
        item["source_varies"] = len(set(item.pop("sources"))) > 1
        item["total_qty"] = round(sum(m["qty"] for m in item["memberships"]
                                      if m["qty"] is not None), 4)
        item["kit_count"] = len(item["memberships"])
    return items


def write_csv(path, header, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


def fmt(value):
    """Format a number for CSV: integers without a decimal tail."""
    if value is None:
        return ""
    if isinstance(value, float):
        value = round(value, 4)
        if value.is_integer():
            return str(int(value))
    return str(value)


def status_label(status):
    return "Current" if status == CURRENT else "Inactive"


def main():
    legacy_rows = read_legacy_rows()
    sp_file, sp_rows = read_sharepoint_rows()
    kits, rows, stats = build(legacy_rows, sp_file, sp_rows)
    items = build_items(rows, kits)

    def sort_key(row):
        return (kits[row["kit_nsn"]]["name"], row["status"] != CURRENT, row["item_nsn"])

    ordered = sorted(rows, key=sort_key)

    # 1. Every kit-item line, the canonical flat table.
    write_csv(
        os.path.join(DATA_DIR, "kit_items.csv"),
        ["Kit NSN", "Kit Name", "Item NSN", "Description", "Quantity", "Unit of Measure",
         "Source", "Unit Price", "Extended Price", "Accountability Code", "Is Sub-Kit",
         "Status"],
        [[r["kit_nsn"], kits[r["kit_nsn"]]["name"], r["item_nsn"], r["description"],
          fmt(r["qty"]), r["uom"], r["source"], fmt(r["price"]), fmt(r["extended_price"]),
          r["accountability_code"], "Y" if r["item_nsn"] in kits else "",
          status_label(r["status"])]
         for r in ordered],
    )

    # 2. One pick list per kit. Remove any left over from a renamed or deleted
    #    kit first: an orphan is indistinguishable from current data.
    os.makedirs(KIT_DIR, exist_ok=True)
    expected = {kit["slug"] + ".csv" for kit in kits.values()}
    for stale in sorted(set(os.listdir(KIT_DIR)) - expected):
        if stale.endswith(".csv"):
            os.remove(os.path.join(KIT_DIR, stale))
            print("removed stale %s" % stale)

    for kit in kits.values():
        kit_lines = [r for r in ordered if r["kit_nsn"] == kit["nsn"]]
        write_csv(
            os.path.join(KIT_DIR, kit["slug"] + ".csv"),
            ["Item NSN", "Description", "Quantity", "Unit of Measure", "Source",
             "Unit Price", "Extended Price", "Accountability Code", "Is Sub-Kit",
             "Status"],
            [[r["item_nsn"], r["description"], fmt(r["qty"]), r["uom"], r["source"],
              fmt(r["price"]), fmt(r["extended_price"]), r["accountability_code"],
              "Y" if r["item_nsn"] in kits else "", status_label(r["status"])]
             for r in kit_lines],
        )

    # 3. Item-level rollup, one row per distinct item across all kits.
    #    "Is Kit" marks containers that have contents of their own. They are
    #    listed because they are real line items inside their parent, but their
    #    value is already the sum of their contents, so anyone counting stock or
    #    money should filter them out — and filter out Status = Inactive too.
    write_csv(
        os.path.join(DATA_DIR, "master_inventory.csv"),
        ["NSN", "Description", "Unit of Measure", "Total Quantity", "Kit Count",
         "Unit Price", "Accountability Code", "Source", "Is Kit", "Kit Membership",
         "Status"],
        [[item["nsn"], item["description"], item["uom"], fmt(item["total_qty"]),
          item["kit_count"], fmt(item["price"]), item["accountability_code"],
          item["source"], "Y" if item["is_kit"] else "",
          "; ".join("%s (%s) x%s" % (kits[m["kit_nsn"]]["name"], m["kit_nsn"],
                                     fmt(m["qty"]))
                    for m in item["memberships"]),
          status_label(item["status"])]
         for item in sorted(items.values(), key=lambda i: i["nsn"])],
    )

    # 4. Structured data for the browser app.
    def count(collection, status):
        return sum(1 for x in collection if x["status"] == status)

    payload = {
        "generated_from": "Data/source/%s (current); Data/legacy/*.xlsx (inactive)"
                          % sp_file,
        "kit_count": len(kits),
        "item_count": len(items),
        "line_count": len(rows),
        "current": {"kits": count(kits.values(), CURRENT),
                    "items": count(items.values(), CURRENT),
                    "lines": count(rows, CURRENT)},
        "inactive": {"kits": count(kits.values(), INACTIVE),
                     "items": count(items.values(), INACTIVE),
                     "lines": count(rows, INACTIVE)},
        "kits": list(kits.values()),
        "items": sorted(items.values(), key=lambda i: i["nsn"]),
    }
    with open(os.path.join(DATA_DIR, "inventory.json"), "w", encoding="utf-8") as handle:
        json.dump(payload, handle, indent=1, ensure_ascii=False)
        handle.write("\n")

    print("current   kits: %(kits)d  items: %(items)d  lines: %(lines)d" % payload["current"])
    print("inactive  kits: %(kits)d  items: %(items)d  lines: %(lines)d" % payload["inactive"])
    if stats["price_fallbacks"]:
        print("previous price kept where SharePoint shows $0: %s"
              % ", ".join(stats["price_fallbacks"]))

    # Keep index.html's asset versions in step, so browsers never pair a cached
    # old app.js with a newer page.
    import stamp_assets
    stamp_assets.main()


if __name__ == "__main__":
    main()
